"""將罪法刑結構化結果另存為「標準化」工作表，提供法條層級與罪名粗分類。"""

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE_SHEET = "罪法刑結構化"
OUTPUT_SHEET = "標準化"
CN_DIGITS = {"零": 0, "〇": 0, "○": 0, "一": 1, "二": 2, "三": 3, "四": 4,
             "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
CN_UNITS = {"十": 10, "百": 100, "千": 1000, "萬": 10_000}
ARTICLE_RE = re.compile(
    r"(?P<name>[^，、；。：（）()\s\[\]]{1,30}?(?:條例|規則|法))"
    r"第(?P<article>[零〇○一二三四五六七八九十百千萬壹貳參叁肆伍陸柒捌玖拾佰仟0-9]+)條"
    r"(?:之(?P<subarticle>[零〇○一二三四五六七八九十百千萬壹貳參叁肆伍陸柒捌玖拾佰仟0-9]+))?"
)
NAME_PREFIX_RE = re.compile(r"^(?:以及|並依|復依|爰依|應依|均依|依照|依|按|適用|違反)+")
NAME_ALIASES = {"中華民國刑法": "刑法", "修正前刑法": "刑法", "修正後刑法": "刑法"}
CHARGE_RULES = (
    (re.compile(r"尿液所含毒品|駕駛動力交通工具.*毒品"), "毒駕罪"),
    (re.compile(r"酒精濃度|酒類.*駕駛|不能安全駕駛"), "酒駕罪"),
    (re.compile(r"過失傷害"), "過失傷害罪"),
    (re.compile(r"過失致死"), "過失致死罪"),
    (re.compile(r"詐欺"), "詐欺罪"),
    (re.compile(r"竊盜"), "竊盜罪"),
    (re.compile(r"施用.*毒品"), "施用毒品罪"),
    (re.compile(r"持有.*毒品"), "持有毒品罪"),
    (re.compile(r"販賣.*毒品"), "販賣毒品罪"),
    (re.compile(r"轉讓.*毒品"), "轉讓毒品罪"),
    (re.compile(r"運輸.*毒品"), "運輸毒品罪"),
    (re.compile(r"製造.*毒品"), "製造毒品罪"),
    (re.compile(r"行使.*文書|偽造.*文書|登載不實"), "偽造文書罪"),
    (re.compile(r"侵占"), "侵占罪"),
    (re.compile(r"恐嚇"), "恐嚇罪"),
    (re.compile(r"毀損"), "毀損罪"),
    (re.compile(r"公然侮辱|誹謗"), "妨害名譽罪"),
    (re.compile(r"傷害"), "傷害罪"),
    (re.compile(r"洗錢"), "洗錢罪"),
    (re.compile(r"犯罪組織"), "組織犯罪罪"),
    (re.compile(r"妨害公務|侮辱公務員"), "妨害公務罪"),
    (re.compile(r"強制性交"), "強制性交罪"),
    (re.compile(r"強制猥褻"), "強制猥褻罪"),
    (re.compile(r"殺人"), "殺人罪"),
    (re.compile(r"強盜"), "強盜罪"),
    (re.compile(r"搶奪"), "搶奪罪"),
    (re.compile(r"強制罪"), "強制罪"),
)


# 將中文或阿拉伯條號轉為整數，供主條及「之N」使用。
def numeral_to_int(value: str) -> int:
    if value.isdigit():
        return int(value)
    value = value.translate(str.maketrans("壹貳參叁肆伍陸柒捌玖拾佰仟", "一二三三四五六七八九十百千"))
    total = section = number = 0
    for char in value:
        if char in CN_DIGITS:
            number = CN_DIGITS[char]
        elif char in CN_UNITS:
            unit = CN_UNITS[char]
            if unit == 10_000:
                total += (section + number) * unit
                section = number = 0
            else:
                section += (number or 1) * unit
                number = 0
    return total + section + number


# 拆開 structure_tasks.py 使用的 [A,B,C] 多標籤格式，空清單不產生項目。
def split_labels(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text or text == "[]":
        return []
    if text.startswith("[") and text.endswith("]"):
        text = text[1:-1]
    return [item.strip() for item in text.split(",") if item.strip()]


# 清除論述前綴與常見別名；同法、同條例由同一列前一個已確認法規承接。
def normalize_law_name(name: str, previous_name: str = "") -> str:
    cleaned = NAME_PREFIX_RE.sub("", name)
    if cleaned in {"同法", "該法", "同條例", "該條例"}:
        return previous_name
    return NAME_ALIASES.get(cleaned, cleaned)


# 將法條同時整理成保留「之N」的標準鍵，以及只到第N條的粗分類。
def normalize_law_levels(value: object) -> tuple[str, str]:
    standard: list[str] = []
    coarse: list[str] = []
    previous_name = ""
    for raw in split_labels(value):
        match = ARTICLE_RE.search(re.sub(r"\s+", "", raw))
        if not match:
            continue
        name = normalize_law_name(match.group("name"), previous_name)
        if not name:
            continue
        previous_name = name
        main = f'{name}第{numeral_to_int(match.group("article"))}條'
        subarticle = match.group("subarticle")
        detailed = f"{main}之{numeral_to_int(subarticle)}" if subarticle else main
        if detailed not in standard:
            standard.append(detailed)
        if main not in coarse:
            coarse.append(main)
    if not standard:
        return "", ""
    return "[" + ",".join(standard) + "]", "[" + ",".join(coarse) + "]"


# 以明確的犯罪家族規則移除細節前綴；任一罪名未命中時整格留空，避免高估成功率。
def normalize_charges(value: object) -> str:
    normalized: list[str] = []
    for charge in split_labels(value):
        standard = next((label for pattern, label in CHARGE_RULES if pattern.search(charge)), "")
        if not standard:
            return ""
        if standard not in normalized:
            normalized.append(standard)
    return "[" + ",".join(normalized) + "]" if normalized else ""


# 讀取結構化工作表並新增標準化工作表；刑期維持原樣以免改變量刑語意。
def normalize_workbook(source: Path, output: Path) -> int:
    if not source.exists():
        raise FileNotFoundError(f"找不到輸入檔：{source}")
    workbook = load_workbook(source)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"找不到工作表：{SOURCE_SHEET}")

    source_sheet = workbook[SOURCE_SHEET]
    headers = [cell.value for cell in source_sheet[1]]
    required = {"裁判字號", "法條", "罪名", "所有宣告刑", "應執行刑"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"缺少必要欄位：{', '.join(sorted(missing))}")

    if OUTPUT_SHEET in workbook.sheetnames:
        del workbook[OUTPUT_SHEET]
    target = workbook.create_sheet(OUTPUT_SHEET)
    target.append([
        "判決書名稱", "法條（原始）", "法條標準鍵", "法條粗分類",
        "罪名（原始）", "罪名（標準化）", "所有宣告刑", "應執行刑",
    ])
    for values in source_sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        standard_laws, coarse_laws = normalize_law_levels(row.get("法條"))
        target.append([
            row.get("裁判字號", ""), row.get("法條", ""), standard_laws, coarse_laws,
            row.get("罪名", ""), normalize_charges(row.get("罪名")),
            row.get("所有宣告刑", ""), row.get("應執行刑", ""),
        ])

    for cell in target[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F3864")
    target.freeze_panes = "A2"
    target.auto_filter.ref = target.dimensions
    for column, width in zip("ABCDEFGH", (38, 45, 45, 45, 45, 35, 32, 32)):
        target.column_dimensions[column].width = width
    for row in target.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output)
    return target.max_row - 1


# 執行最小自我檢查，涵蓋兩層法條、別名承接及常見罪名粗分類。
def self_check() -> None:
    assert normalize_law_levels("[刑法第一百八十五條之三第1項,同法第339條之4]") == (
        "[刑法第185條之3,刑法第339條之4]", "[刑法第185條,刑法第339條]"
    )
    assert normalize_law_levels(
        "[中華民國刑法第參佰貳拾條第1項,修正前刑法第320條]"
    ) == (
        "[刑法第320條]", "[刑法第320條]"
    )
    assert normalize_law_levels("[]") == ("", "")
    assert normalize_charges("[三人以上共同詐欺取財未遂罪,攜帶兇器竊盜罪]") == (
        "[詐欺罪,竊盜罪]"
    )
    assert normalize_charges("[吐氣所含酒精濃度達每公升0.25毫克以上而駕駛動力交通工具罪]") == (
        "[酒駕罪]"
    )
    assert normalize_charges("[駕駛動力交通工具而有尿液所含毒品達公告濃度值以上之情形罪]") == (
        "[毒駕罪]"
    )
    assert normalize_charges("[]") == ""
    assert normalize_charges("[尚未建立分類之罪]") == ""


# 提供原檔新增工作表的預設操作，也允許另存新檔以保留輸入檔。
def main() -> None:
    parser = argparse.ArgumentParser(description="新增罪法刑標準化工作表")
    parser.add_argument("input", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.input
    self_check()
    count = normalize_workbook(args.input, output)
    print(f"完成：{count} 筆 → {output}（{OUTPUT_SHEET}）")


if __name__ == "__main__":
    main()
