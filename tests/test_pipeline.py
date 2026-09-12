# -*- coding: utf-8 -*-
"""
pipeline.py 的離線測試：以 --skip-crawl 路徑驗證「標籤 → 匯出篩選 → 檔名」這條串接。
不連網、不開瀏覽器，全部在暫存 DB 與暫存目錄進行。
"""

import os
import sqlite3
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import export_excel  # noqa: E402
import html_parser  # noqa: E402
import pipeline  # noqa: E402
from crawl_batched import LABEL_ALL, build_label  # noqa: E402

ROWS = [
    ("臺灣臺北地方法院 113 年度訴字第 1 號民事判決", "臺灣臺北地方法院",
     "2025-01-07", "判決", "臺灣臺北地方法院-民事-判決"),
    ("臺灣高雄地方法院 112 年度訴字第 2 號民事判決", "臺灣高雄地方法院",
     "2023-06-01", "判決", "借名登記"),
]


class TestPipelineSkipCrawl(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        db = os.path.join(self.tmp.name, "t.db")
        self._orig = (html_parser.DB_PATH, export_excel.DB_PATH)
        html_parser.DB_PATH = export_excel.DB_PATH = db
        html_parser.init_db()
        conn = sqlite3.connect(db)
        conn.executemany(
            """INSERT INTO judgments
               (crawl_id, case_number, court, judgment_date, judgment_type, keyword)
               VALUES (?,?,?,?,?,?)""",
            [(i, *r) for i, r in enumerate(ROWS, start=1)],
        )
        conn.commit()
        conn.close()

    def tearDown(self):
        html_parser.DB_PATH, export_excel.DB_PATH = self._orig
        self.tmp.cleanup()

    def _run(self, out_name, **kwargs):
        out = os.path.join(self.tmp.name, out_name)
        pipeline.run(
            keyword=kwargs.pop("keyword", ""),
            max_results=0, headless=True, output=out, full_text=False,
            start_date=date(2025, 1, 1), end_date=date(2025, 12, 31),
            skip_crawl=True, **kwargs,
        )
        return out

    def _row_count(self, xlsx):
        from openpyxl import load_workbook
        return load_workbook(xlsx).active.max_row - 1      # 扣掉標題列

    def test_condition_label_filters_export(self):
        out = self._run("cond.xlsx", court="臺灣臺北地方法院",
                        case_types=("民事",), judgment_type="判決")
        self.assertEqual(self._row_count(out), 1)

    def test_keyword_filters_export(self):
        out = self._run("kw.xlsx", keyword="借名登記")
        self.assertEqual(self._row_count(out), 1)

    def test_no_condition_exports_everything(self):
        # 只下 --skip-crawl 時不應該把「全部」當成關鍵字去比對，否則會匯出 0 筆
        self.assertEqual(build_label(""), LABEL_ALL)
        out = self._run("all.xlsx")
        self.assertEqual(self._row_count(out), len(ROWS))

    def test_default_filename_uses_label(self):
        # 未指定 -o 時，檔名以標籤命名（標籤含空白等字元會被換成底線）
        cwd = os.getcwd()
        os.chdir(self.tmp.name)
        try:
            pipeline.run(
                keyword="", max_results=0, headless=True, output="", full_text=False,
                start_date=date(2025, 1, 1), end_date=date(2025, 12, 31),
                skip_crawl=True, court="臺灣臺北地方法院",
                case_types=("民事",), judgment_type="判決",
            )
            produced = [f for f in os.listdir(".") if f.endswith(".xlsx")]
        finally:
            os.chdir(cwd)
        self.assertEqual(len(produced), 1)
        self.assertTrue(produced[0].startswith("judgments_臺灣臺北地方法院-民事-判決_"),
                        msg=produced[0])


if __name__ == "__main__":
    unittest.main()
