# -*- coding: utf-8 -*-
"""
export_excel.py 的離線測試：篩選條件（關鍵字／法院／裁判日期）與 Excel 產出。
全部在暫存 DB 與暫存目錄進行，不碰真正的 judgments.db。
"""

import os
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import export_excel  # noqa: E402
import html_parser  # noqa: E402
from export_excel import BASE_COLUMNS, export_to_excel, fetch_judgments  # noqa: E402

ROWS = [
    # (case_number, court, judgment_date, judgment_type, keyword, verdict)
    ("臺灣臺北地方法院 113 年度訴字第 6910 號民事判決", "臺灣臺北地方法院",
     "2025-01-07", "判決", "臺灣臺北地方法院-民事-判決", "被告應給付原告…"),
    ("臺灣臺北地方法院 113 年度除字第 1836 號民事判決", "臺灣臺北地方法院",
     "2025-01-08", "判決", "臺灣臺北地方法院-民事-判決", "宣告…證券無效"),
    ("臺灣高雄地方法院 112 年度訴字第 1 號民事判決", "臺灣高雄地方法院",
     "2023-06-01", "判決", "借名登記", "原告之訴駁回"),
]


class _TempDbCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        db = os.path.join(self.tmp.name, "t.db")
        self._orig = (export_excel.DB_PATH, html_parser.DB_PATH)
        export_excel.DB_PATH = html_parser.DB_PATH = db
        html_parser.init_db()
        conn = sqlite3.connect(db)
        conn.executemany(
            """INSERT INTO judgments
               (crawl_id, case_number, court, judgment_date, judgment_type,
                keyword, verdict)
               VALUES (?,?,?,?,?,?,?)""",
            [(i, *r) for i, r in enumerate(ROWS, start=1)],
        )
        conn.commit()
        conn.close()

    def tearDown(self):
        export_excel.DB_PATH, html_parser.DB_PATH = self._orig
        self.tmp.cleanup()


class TestFetchJudgments(_TempDbCase):
    def test_no_filter_returns_all(self):
        self.assertEqual(len(fetch_judgments(limit=0)), 3)

    def test_limit(self):
        self.assertEqual(len(fetch_judgments(limit=2)), 2)

    def test_keyword_matches_label(self):
        rows = fetch_judgments(limit=0, keyword="臺灣臺北地方法院-民事-判決")
        self.assertEqual(len(rows), 2)

    def test_keyword_also_matches_case_number_and_verdict(self):
        # -k 同時比對 keyword／裁判字號／主文，維持既有行為
        self.assertEqual(len(fetch_judgments(limit=0, keyword="除字")), 1)
        self.assertEqual(len(fetch_judgments(limit=0, keyword="駁回")), 1)

    def test_court_filter_is_partial_match(self):
        self.assertEqual(len(fetch_judgments(limit=0, court="臺灣臺北")), 2)
        self.assertEqual(len(fetch_judgments(limit=0, court="臺灣高雄地方法院")), 1)

    def test_date_filters_accept_roc_and_ad(self):
        # judgment_date 以西元 ISO 儲存，輸入民國格式也會先正規化
        self.assertEqual(len(fetch_judgments(limit=0, start_date="2025/01/01")), 2)
        self.assertEqual(len(fetch_judgments(limit=0, start_date="114/01/08")), 1)
        self.assertEqual(len(fetch_judgments(limit=0, end_date="2024/12/31")), 1)

    def test_combined_filters(self):
        rows = fetch_judgments(limit=0, court="臺灣臺北",
                               start_date="2025/01/08", end_date="2025/01/08")
        self.assertEqual(len(rows), 1)
        self.assertIn("除字", rows[0]["case_number"])


class TestExportToExcel(_TempDbCase):
    def test_creates_file_with_expected_headers(self):
        from openpyxl import load_workbook
        out = os.path.join(self.tmp.name, "out.xlsx")
        rows = fetch_judgments(limit=0)
        self.assertTrue(export_to_excel(rows, out, include_full_text=False))
        self.assertTrue(os.path.exists(out))

        ws = load_workbook(out).active
        headers = [c.value for c in ws[1]]
        self.assertIn("裁判字號", headers)
        self.assertIn("裁判種類", headers)
        self.assertIn("搜尋關鍵字", headers)
        self.assertNotIn("全文", headers)          # 未加 --full-text
        self.assertEqual(ws.max_row, len(rows) + 1)

    def test_full_text_column_added(self):
        from openpyxl import load_workbook
        out = os.path.join(self.tmp.name, "full.xlsx")
        export_to_excel(fetch_judgments(limit=0), out, include_full_text=True)
        headers = [c.value for c in load_workbook(out).active[1]]
        self.assertIn("全文", headers)

    def test_base_columns_cover_db_schema(self):
        # BASE_COLUMNS 的每個欄位都必須存在於 judgments 表，否則匯出會整批失敗
        conn = sqlite3.connect(export_excel.DB_PATH)
        cols = {r[1] for r in conn.execute("PRAGMA table_info(judgments)")}
        conn.close()
        for db_col, _label in BASE_COLUMNS:
            self.assertIn(db_col, cols, msg=db_col)


if __name__ == "__main__":
    unittest.main()
